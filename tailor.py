"""Cover letter tailoring helper.

The cover letter template is a Word document whose bold runs are placeholders.
Claude (or a person) reads `context`, decides what goes in each slot, writes a
slots JSON file, and runs `render` to produce the tailored .docx and .pdf.

Usage:
    python tailor.py list
    python tailor.py fetch <application_id>                 # save job description from the posting URL
    python tailor.py set-description <application_id> <file>
    python tailor.py context <application_id>               # everything needed to fill the slots
    python tailor.py render <application_id> <slots.json>

slots.json:
    {"slots": [{"text": "...", "before": "optional replacement for the run just before the slot"}, ...]}
One entry per placeholder, in template order. Slot 1 may be omitted text ("") to use today's date
if the placeholder is DATE.
"""
import json
import os
import re
import subprocess
import sys
import urllib.request
from datetime import date
from pathlib import Path

import docx
import pdfplumber
import pymysql
from bs4 import BeautifulSoup
from dotenv import load_dotenv

load_dotenv(Path(__file__).with_name(".env"))

TEMPLATE = os.getenv("COVER_LETTER_TEMPLATE")
RESUME = os.getenv("RESUME_DOCX")
TAILORED_DIR = os.getenv("TAILORED_DIR")
MAX_PAGES = int(os.getenv("COVER_LETTER_MAX_PAGES", "1"))


def db():
    return pymysql.connect(
        host=os.getenv("DB_HOST", "127.0.0.1"),
        port=int(os.getenv("DB_PORT", "3306")),
        user=os.getenv("DB_USER", "jobtracker"),
        password=os.getenv("DB_PASSWORD", ""),
        database=os.getenv("DB_NAME", "job_tracker"),
        cursorclass=pymysql.cursors.DictCursor,
        autocommit=True,
    )


def get_application(app_id):
    with db().cursor() as cur:
        cur.execute("SELECT * FROM applications WHERE id = %s", (app_id,))
        row = cur.fetchone()
    if row is None:
        sys.exit(f"No application with id {app_id}")
    return row


def find_slots(document):
    """Return [(paragraph, run_index)] for every bold, non-empty run, in document order."""
    slots = []
    for p in document.paragraphs:
        for i, r in enumerate(p.runs):
            if r.bold and r.text.strip():
                slots.append((p, i))
    return slots


def previous_text_run(paragraph, index):
    for j in range(index - 1, -1, -1):
        if paragraph.runs[j].text:
            return j
    return None


def docx_text(path):
    return "\n".join(p.text for p in docx.Document(path).paragraphs if p.text.strip())


def cmd_list():
    with db().cursor() as cur:
        cur.execute("""
            SELECT a.id, a.company, a.role, a.status, a.deadline,
                   a.job_description IS NOT NULL AS has_jd,
                   (SELECT COUNT(*) FROM materials m WHERE m.application_id = a.id) AS letters
            FROM applications a ORDER BY a.id""")
        for r in cur.fetchall():
            print(f"{r['id']:>3}  {r['status']:<10} JD:{'yes' if r['has_jd'] else 'no ':<3}  "
                  f"letters:{r['letters']}  deadline:{str(r['deadline'] or '-'):<10}  {r['company']} | {r['role']}")


def save_description(app_id, text):
    text = text.strip()
    with db().cursor() as cur:
        cur.execute("UPDATE applications SET job_description = %s WHERE id = %s", (text, app_id))
    print(f"Saved {len(text):,} characters to application {app_id}")


def cmd_fetch(app_id):
    url = get_application(app_id)["url"]
    if not url:
        sys.exit("Application has no URL")
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 (job-tracker)"})
    with urllib.request.urlopen(req, timeout=20) as resp:
        soup = BeautifulSoup(resp.read(), "html.parser")
    for tag in soup(["script", "style", "noscript", "header", "footer", "nav", "svg"]):
        tag.decompose()
    text = re.sub(r"\n\s*\n+", "\n\n", soup.get_text("\n")).strip()
    if len(text) < 500:
        sys.exit(f"Only got {len(text)} characters; the page probably needs a browser. Paste it instead.")
    save_description(app_id, text)


def cmd_context(app_id):
    a = get_application(app_id)
    print("=== APPLICATION ===")
    for k in ["id", "company", "role", "url", "location", "work_mode", "salary_range", "deadline", "notes"]:
        print(f"{k}: {a[k]}")
    print("\n=== JOB DESCRIPTION ===")
    print(a["job_description"] or "(none saved - run `fetch` or `set-description` first)")
    print("\n=== RESUME (facts may only come from here) ===")
    print(docx_text(RESUME))
    print("\n=== COVER LETTER TEMPLATE ===")
    print(docx_text(TEMPLATE))
    print("\n=== SLOTS (fill in this order) ===")
    for n, (p, i) in enumerate(find_slots(docx.Document(TEMPLATE)), 1):
        j = previous_text_run(p, i)
        before = repr(p.runs[j].text) if j is not None else "(none)"
        after = "".join(r.text for r in p.runs[i + 1:])[:60]
        print(f"{n}. {p.runs[i].text.strip()!r:<45} run before: {before:<18} then: {after!r}")


def safe_name(s, limit=60):
    s = re.sub(r"[^\w\s&-]", "", s).strip()
    return re.sub(r"\s+", "_", s)[:limit].rstrip("_")


def to_windows(path):
    return subprocess.run(["wslpath", "-w", str(path)], capture_output=True, text=True, check=True).stdout.strip()


def word_to_pdf(docx_path, pdf_path):
    script = (
        "$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
        f"$d = $w.Documents.Open('{to_windows(docx_path)}', $false, $true); "
        f"$d.SaveAs2('{to_windows(pdf_path)}', 17); $d.Close($false); $w.Quit()"
    )
    subprocess.run(["powershell.exe", "-NoProfile", "-Command", script], check=True, capture_output=True)


def cmd_render(app_id, slots_file):
    a = get_application(app_id)
    values = json.loads(Path(slots_file).read_text())["slots"]
    document = docx.Document(TEMPLATE)
    slots = find_slots(document)
    if len(values) != len(slots):
        sys.exit(f"Template has {len(slots)} slots but {len(values)} values were given")

    edits = []
    for (p, i), v in zip(slots, values):
        run = p.runs[i]
        placeholder = run.text
        text = v["text"] or (date.today().strftime("%B %-d, %Y") if placeholder.strip() == "DATE" else "")
        if not text:
            sys.exit(f"Empty value for slot {placeholder!r}")
        # keep the placeholder's trailing whitespace so surrounding spacing is unchanged
        run.text = text + placeholder[len(placeholder.rstrip()):]
        run.bold = False
        if "before" in v:
            j = previous_text_run(p, i)
            if j is None:
                sys.exit(f"Slot {placeholder!r} has no run before it to change")
            edits.append(f"{p.runs[j].text!r} -> {v['before']!r} (before {placeholder.strip()})")
            p.runs[j].text = v["before"]

    out_dir = Path(TAILORED_DIR) / safe_name(a["company"])
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = f"{Path(TEMPLATE).stem}_{safe_name(a['role'], 50)}"
    docx_path, pdf_path = out_dir / f"{stem}.docx", out_dir / f"{stem}.pdf"
    document.save(docx_path)
    word_to_pdf(docx_path, pdf_path)

    with pdfplumber.open(pdf_path) as pdf:
        pages = len(pdf.pages)
    if pages > MAX_PAGES:
        docx_path.unlink()
        pdf_path.unlink()
        sys.exit(f"REJECTED: letter runs to {pages} pages (limit {MAX_PAGES}). Shorten the slot text and render again.")

    with db().cursor() as cur:
        # re-rendering the same letter replaces its previous record
        cur.execute("DELETE FROM materials WHERE application_id = %s AND pdf_path = %s", (app_id, str(pdf_path)))
        cur.execute(
            "INSERT INTO materials (application_id, kind, docx_path, pdf_path, slot_values) VALUES (%s, %s, %s, %s, %s)",
            (app_id, "cover_letter", str(docx_path), str(pdf_path), json.dumps(values)),
        )
    print("=== RENDERED LETTER ===")
    print(docx_text(docx_path))
    print("\n=== CHANGES OUTSIDE SLOTS ===")
    print("\n".join(edits) or "(none)")
    print(f"\nSaved:\n  {docx_path}\n  {pdf_path}")


def main():
    args = sys.argv[1:]
    commands = {
        "list": (cmd_list, 0),
        "fetch": (cmd_fetch, 1),
        "set-description": (lambda i, f: save_description(i, Path(f).read_text()), 2),
        "context": (cmd_context, 1),
        "render": (cmd_render, 2),
    }
    if not args or args[0] not in commands or len(args) - 1 != commands[args[0]][1]:
        sys.exit(__doc__)
    fn, _ = commands[args[0]]
    fn(*[int(x) if x.isdigit() else x for x in args[1:]])


if __name__ == "__main__":
    main()
